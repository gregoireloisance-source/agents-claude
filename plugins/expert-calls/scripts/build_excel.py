#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_excel.py : génère le classeur Excel d'une recherche d'experts.

Usage : python3 build_excel.py donnees.json sortie.xlsx
Dépendance : openpyxl (pip install openpyxl)
"""
import json
import math
import sys
from urllib.parse import quote, urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

FONT = "Arial"
NAVY = "1F3864"
BAND = "F2F2F2"
LINE = "BFBFBF"
MAX_ROWS = 200  # lignes couvertes par les listes déroulantes et les formules

STATUTS = [
    "À contacter",
    "Message envoyé",
    "Réponse reçue",
    "Call planifié",
    "Call réalisé",
    "Décliné / sans réponse",
]

LEVEL_FILLS = {
    "conf": {"haute": "C6E0B4", "moyenne": "FFE699", "faible": "F8CBAD"},
    "prio": {"haute": "9DC3E6", "moyenne": "DDEBF7", "basse": "FFFFFF"},
    "cov": {"bonne": "C6E0B4", "partielle": "FFE699", "insuffisante": "F8CBAD"},
}

THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(bold=False, color="000000", size=10, italic=False, underline=None):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic, underline=underline)


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def domain(url):
    host = urlparse(url).netloc.lower()
    if host.startswith("www."):
        host = host[4:]
    return host or url


def est_lines(text, width):
    """Estimation du nombre de lignes d'une cellule avec retour à la ligne."""
    if text is None:
        return 1
    per_line = max(int(width * 1.1), 8)
    total = 0
    for part in str(text).split("\n"):
        total += max(1, math.ceil(len(part) / per_line))
    return total


def norm_sources(items):
    """Accepte une liste de chaînes (URL) ou de dicts {url, note}."""
    out = []
    for item in items or []:
        if isinstance(item, dict):
            url, note = item.get("url", ""), item.get("note", "")
        else:
            url, note = str(item), ""
        if url:
            out.append((url, note))
    return out


def link_cell(sources, index):
    if index < len(sources):
        url = sources[index][0]
        return (url, domain(url))
    return None


def search_links(name_or_function, company):
    """Liens de recherche LinkedIn et Google préremplis (simple ouverture dans le navigateur)."""
    terms = " ".join(t for t in (name_or_function, company) if t).strip()
    if not terms:
        return None, None
    linkedin = "https://www.linkedin.com/search/results/people/?keywords=" + quote(terms)
    quoted = " ".join(f'"{t}"' for t in (name_or_function, company) if t)
    google = "https://www.google.com/search?q=" + quote(f"site:linkedin.com/in {quoted}")
    return (linkedin, "Rechercher"), (google, "Rechercher")


def page_setup(ws, header_row=None):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if header_row:
        ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.6


def build_sheet(wb, title, heading, subtitle, columns, rows, freeze_cols=0, statut_col=None):
    """
    columns : liste de (en-tête, largeur, type)
        types : text, center, prio, conf, cov, link, linkfull
    rows : liste de listes de valeurs ; pour link/linkfull : (url, libellé) ou None
    """
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90

    ws["A1"] = heading
    ws["A1"].font = font(True, NAVY, 14)
    ws["A2"] = subtitle
    ws["A2"].font = font(False, "7F7F7F", 9, italic=True)

    header_row = 4
    for j, (name, width, _kind) in enumerate(columns, start=1):
        c = ws.cell(row=header_row, column=j, value=name)
        c.font = font(True, "FFFFFF")
        c.fill = fill(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[header_row].height = 32

    for i, row in enumerate(rows):
        r = header_row + 1 + i
        lines = 1
        for j, (_name, width, kind) in enumerate(columns, start=1):
            value = row[j - 1] if j - 1 < len(row) else None
            c = ws.cell(row=r, column=j)
            c.border = BORDER
            c.font = font()
            centered = kind in ("center", "prio", "conf", "cov")
            c.alignment = Alignment(
                vertical="top", wrap_text=True, horizontal="center" if centered else "left"
            )
            if i % 2 == 1:
                c.fill = fill(BAND)
            if kind in ("link", "linkfull"):
                if value:
                    url, label = value
                    c.value = url if kind == "linkfull" else label
                    if url.lower().startswith("http"):
                        c.hyperlink = url
                        c.font = font(color="0563C1", underline="single")
                    lines = max(lines, est_lines(c.value, width) if kind == "linkfull" else 1)
            else:
                c.value = value
                lines = max(lines, est_lines(value, width))
                if kind in LEVEL_FILLS and value:
                    color = LEVEL_FILLS[kind].get(str(value).strip().lower())
                    if color:
                        c.fill = fill(color)
        ws.row_dimensions[r].height = min(max(16, 13 * lines + 5), 400)

    last = header_row + max(len(rows), 1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{last}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=freeze_cols + 1)

    if statut_col:
        letter = get_column_letter(statut_col)
        formula = '"' + ",".join(STATUTS) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.error = "Choisis une valeur de la liste."
        dv.errorTitle = "Statut"
        ws.add_data_validation(dv)
        dv.add(f"{letter}{header_row + 1}:{letter}{header_row + MAX_ROWS}")

    page_setup(ws, header_row)
    return ws


def section_header(ws, row, text, ncols=2):
    for j in range(1, ncols + 1):
        ws.cell(row=row, column=j).fill = fill(NAVY)
    c = ws.cell(row=row, column=1, value=text)
    c.font = font(True, "FFFFFF")
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 20


def kv_row(ws, row, label, value, link=False, width_b=95):
    a = ws.cell(row=row, column=1, value=label)
    a.font = font(True)
    a.alignment = Alignment(vertical="top", wrap_text=True)
    a.border = BORDER
    a.fill = fill("D9E1F2")
    b = ws.cell(row=row, column=2, value=value)
    b.font = font()
    b.alignment = Alignment(vertical="top", wrap_text=True, horizontal="left")
    b.border = BORDER
    if link and value and str(value).lower().startswith("http"):
        b.hyperlink = str(value)
        b.font = font(color="0563C1", underline="single")
    ws.row_dimensions[row].height = max(16, 13 * est_lines(value, width_b) + 5)


def build_summary(wb, data, col_a, col_b):
    ws = wb.active
    ws.title = "Synthèse"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 95

    target = data.get("target", "Société cible")
    ws["A1"] = f"Recherche d'experts : {target}"
    ws["A1"].font = font(True, NAVY, 16)
    ws["A2"] = (
        "Document de travail confidentiel. Sources publiques uniquement. "
        f"Généré le {data.get('date', '')}."
    )
    ws["A2"].font = font(False, "7F7F7F", 9, italic=True)

    r = 4
    section_header(ws, r, "Périmètre")
    items = [
        ("Société cible", target, False),
        ("Site web", data.get("website", ""), True),
        ("Zone géographique", data.get("geography", ""), False),
        ("Marchés et segments", data.get("scope", ""), False),
        ("Concurrents", data.get("competitors", ""), False),
        ("Partenaires", data.get("partners", ""), False),
    ]
    for label, value, is_link in items:
        r += 1
        kv_row(ws, r, label, value, link=is_link)

    r += 2
    section_header(ws, r, "Suivi du pipeline (se met à jour avec la colonne Statut)")
    sheet_a, col_a_letter = "'Experts A'", get_column_letter(col_a)
    sheet_b, col_b_letter = "'Pistes B'", get_column_letter(col_b)
    first_status_row = r + 1
    for statut in STATUTS:
        r += 1
        ws.cell(row=r, column=1, value=statut).font = font()
        formula = (
            f"=COUNTIF({sheet_a}!${col_a_letter}$5:${col_a_letter}${MAX_ROWS + 4},A{r})"
            f"+COUNTIF({sheet_b}!${col_b_letter}$5:${col_b_letter}${MAX_ROWS + 4},A{r})"
        )
        v = ws.cell(row=r, column=2, value=formula)
        v.font = font()
        v.alignment = Alignment(horizontal="left")
        for j in (1, 2):
            ws.cell(row=r, column=j).border = BORDER
    last_status_row = r
    r += 1
    ws.cell(row=r, column=1, value="Total des contacts").font = font(True)
    t = ws.cell(row=r, column=2, value=f"=SUM(B{first_status_row}:B{last_status_row})")
    t.font = font(True)
    t.alignment = Alignment(horizontal="left")
    for j in (1, 2):
        ws.cell(row=r, column=j).border = BORDER
        ws.cell(row=r, column=j).fill = fill("D9E1F2")
    r += 1
    ws.cell(row=r, column=1, value="dont experts nommés (niveau A)").font = font()
    ws.cell(row=r, column=2, value=f"=COUNTA({sheet_a}!$A$5:$A${MAX_ROWS + 4})").font = font()
    ws.cell(row=r, column=2).alignment = Alignment(horizontal="left")
    r += 1
    ws.cell(row=r, column=1, value="dont pistes sans nom (niveau B)").font = font()
    ws.cell(row=r, column=2, value=f"=COUNTA({sheet_b}!$A$5:$A${MAX_ROWS + 4})").font = font()
    ws.cell(row=r, column=2).alignment = Alignment(horizontal="left")
    for rr in (r - 1, r):
        for j in (1, 2):
            ws.cell(row=rr, column=j).border = BORDER

    r += 2
    section_header(ws, r, "Lecture du document")
    legend = [
        ("Niveau A", "Experts nommés, avec au moins une source publique vérifiable."),
        ("Niveau B", "Pistes : entreprise et fonction à rechercher, sans personne identifiée."),
        ("Confiance haute", "Deux sources publiques indépendantes confirment le poste et le lien avec le sujet."),
        ("Confiance moyenne", "Une seule source publique."),
        ("Confiance faible", "Lien déduit ou non vérifié."),
        ("Liens de recherche", "Les liens LinkedIn et Google ouvrent une recherche dans ton navigateur. Ils ne vérifient rien : le tri et la vérification des profils restent manuels."),
        ("Affirmation de l'entreprise", "Un client cité sur le site d'une société reste une affirmation de cette société, pas un fait vérifié."),
    ]
    for label, text in legend:
        r += 1
        kv_row(ws, r, label, text)

    r += 2
    section_header(ws, r, "Précautions")
    cautions = [
        "Ne contacter aucun salarié actuel de la cible. Ne demander aucune information confidentielle de l'employeur actuel d'un expert.",
        "Valider l'approche avec la compliance avant tout envoi (informations privilégiées, RGPD, politique interne).",
        "Aucun message n'a été envoyé : tous les textes sont des brouillons à relire et à signer.",
    ]
    for text in cautions:
        r += 1
        ws.cell(row=r, column=1, value="•").alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=r, column=1).font = font()
        c = ws.cell(row=r, column=2, value=text)
        c.font = font()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = max(16, 13 * est_lines(text, 95) + 5)

    page_setup(ws)
    return ws


def guide_rows(ref, label, guide):
    rows = []
    n = 0
    for item in guide or []:
        block = item.get("block", "")
        question = item.get("question", "")
        is_question = not block.lower().startswith(("introduction", "clôture", "cloture", "à faire", "a faire", "prudence"))
        if is_question:
            n += 1
        rows.append(
            [
                ref,
                label,
                block,
                n if is_question else "",
                question,
                item.get("followups", ""),
                item.get("serves", ""),
            ]
        )
    return rows


def build_workbook(data):
    wb = Workbook()
    target = data.get("target", "Société cible")
    subtitle = f"{target} | généré le {data.get('date', '')} | sources publiques uniquement"

    experts = data.get("experts_a", [])
    leads = data.get("leads_b", [])
    templates = data.get("templates_b", [])
    questions = data.get("questions", [])

    # Colonnes de statut (pour les formules de la synthèse)
    cols_a = [
        ("Réf.", 7, "center"),
        ("Nom", 24, "text"),
        ("Fonction", 30, "text"),
        ("Entreprise", 24, "text"),
        ("Catégorie", 22, "text"),
        ("Pertinence", 52, "text"),
        ("Questions couvertes", 13, "center"),
        ("Priorité", 11, "prio"),
        ("Confiance", 11, "conf"),
        ("Source 1", 22, "link"),
        ("Source 2", 22, "link"),
        ("Recherche LinkedIn", 14, "link"),
        ("Recherche Google", 14, "link"),
        ("Statut", 22, "center"),
        ("Commentaires", 32, "text"),
    ]
    cols_b = [
        ("Réf.", 7, "center"),
        ("Type de profil", 24, "text"),
        ("Entreprise", 26, "text"),
        ("Fonction visée", 28, "text"),
        ("Raison", 48, "text"),
        ("Questions couvertes", 13, "center"),
        ("Mots-clés de recherche", 38, "text"),
        ("Source", 22, "link"),
        ("Recherche LinkedIn", 14, "link"),
        ("Recherche Google", 14, "link"),
        ("Statut", 22, "center"),
        ("Commentaires", 32, "text"),
    ]
    statut_a = [i for i, c in enumerate(cols_a, 1) if c[0] == "Statut"][0]
    statut_b = [i for i, c in enumerate(cols_b, 1) if c[0] == "Statut"][0]

    build_summary(wb, data, statut_a, statut_b)

    # Questions et couverture
    q_rows = [
        [
            q.get("id", i + 1),
            q.get("question", ""),
            q.get("why", ""),
            q.get("expert_types", ""),
            q.get("experts", ""),
            q.get("coverage", ""),
            q.get("gap", ""),
        ]
        for i, q in enumerate(questions)
    ]
    build_sheet(
        wb,
        "Questions",
        "Questions d'investissement et couverture",
        subtitle,
        [
            ("N°", 6, "center"),
            ("Question d'investissement", 52, "text"),
            ("Pourquoi elle compte", 42, "text"),
            ("Types d'experts utiles", 32, "text"),
            ("Experts identifiés (réf.)", 22, "text"),
            ("Couverture", 16, "cov"),
            ("Manques et profils complémentaires", 46, "text"),
        ],
        q_rows,
        freeze_cols=2,
    )

    # Experts nommés
    a_rows = []
    for e in experts:
        src = norm_sources(e.get("sources"))
        li, go = search_links(e.get("name", ""), e.get("company", ""))
        a_rows.append(
            [
                e.get("id", ""),
                e.get("name", ""),
                e.get("role", ""),
                e.get("company", ""),
                e.get("category", ""),
                e.get("relevance", ""),
                e.get("questions", ""),
                e.get("priority", ""),
                e.get("confidence", ""),
                link_cell(src, 0),
                link_cell(src, 1),
                li,
                go,
                "À contacter",
                "",
            ]
        )
    build_sheet(
        wb, "Experts A", "Niveau A : experts nommés et sourcés", subtitle,
        cols_a, a_rows, freeze_cols=2, statut_col=statut_a,
    )

    # Pistes
    b_rows = []
    for lead in leads:
        src = norm_sources(lead.get("sources"))
        li, go = search_links(lead.get("function", ""), lead.get("company", ""))
        b_rows.append(
            [
                lead.get("id", ""),
                lead.get("profile_type", ""),
                lead.get("company", ""),
                lead.get("function", ""),
                lead.get("reason", ""),
                lead.get("questions", ""),
                lead.get("keywords", ""),
                link_cell(src, 0),
                li,
                go,
                "À contacter",
                "",
            ]
        )
    build_sheet(
        wb, "Pistes B", "Niveau B : pistes à identifier (aucun nom confirmé)", subtitle,
        cols_b, b_rows, freeze_cols=2, statut_col=statut_b,
    )

    # Messages d'approche
    m_rows = []
    for e in experts:
        m = e.get("message", {}) or {}
        m_rows.append(
            [
                e.get("id", ""),
                f"{e.get('name', '')} ({e.get('company', '')})",
                m.get("language", ""),
                m.get("subject", ""),
                m.get("body", ""),
            ]
        )
    for t in templates:
        m_rows.append(
            [
                "Modèle",
                t.get("profile_type", ""),
                t.get("language", ""),
                t.get("subject", ""),
                t.get("message", ""),
            ]
        )
    build_sheet(
        wb,
        "Messages",
        "Brouillons de messages d'approche",
        subtitle + " | brouillons à relire, aucun envoi automatique",
        [
            ("Réf.", 9, "center"),
            ("Destinataire / type de profil", 30, "text"),
            ("Langue", 9, "center"),
            ("Objet", 38, "text"),
            ("Message", 105, "text"),
        ],
        m_rows,
        freeze_cols=2,
    )

    # Guides d'entretien
    g_rows = []
    for e in experts:
        g_rows += guide_rows(e.get("id", ""), f"{e.get('name', '')} ({e.get('company', '')})", e.get("guide"))
    for t in templates:
        g_rows += guide_rows("Modèle", t.get("profile_type", ""), t.get("guide"))
    build_sheet(
        wb,
        "Guides",
        "Guides d'entretien (45 minutes)",
        subtitle,
        [
            ("Réf.", 9, "center"),
            ("Expert / type de profil", 28, "text"),
            ("Bloc", 26, "text"),
            ("N°", 5, "center"),
            ("Question", 62, "text"),
            ("Relances", 52, "text"),
            ("Questions d'inv.", 12, "center"),
        ],
        g_rows,
        freeze_cols=2,
    )

    # Sources
    s_rows = []
    for e in experts:
        for url, note in norm_sources(e.get("sources")):
            s_rows.append([e.get("id", ""), f"{e.get('name', '')} ({e.get('company', '')})", note, (url, url)])
    for lead in leads:
        for url, note in norm_sources(lead.get("sources")):
            s_rows.append([lead.get("id", ""), lead.get("company", ""), note, (url, url)])
    for url, note in norm_sources(data.get("market_sources")):
        s_rows.append(["Marché", "Cible et marché", note, (url, url)])
    build_sheet(
        wb,
        "Sources",
        "Sources publiques",
        subtitle,
        [
            ("Réf.", 9, "center"),
            ("Expert / entreprise", 34, "text"),
            ("Ce que la source établit", 62, "text"),
            ("Lien", 70, "linkfull"),
        ],
        s_rows,
        freeze_cols=2,
    )

    return wb


def main():
    if len(sys.argv) != 3:
        print("Usage : python3 build_excel.py donnees.json sortie.xlsx")
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        data = json.load(f)
    wb = build_workbook(data)
    wb.save(sys.argv[2])
    print("Fichier créé :", sys.argv[2])


if __name__ == "__main__":
    main()